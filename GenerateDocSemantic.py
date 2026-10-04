import os
import re
from docx import Document
from docx.shared import Pt, RGBColor

station_dir = r"h:\水晶光电AOI程序1002\Autoframe\Station"
output_file = r"h:\水晶光电AOI程序1002\Station_Workflows_Vietnamese.docx"

word_dict = {
    "上料站": "Trạm Cấp Liệu", "拉料站": "Trạm Kéo Liệu", "扩膜站": "Trạm Căng Màng",
    "摆臂站": "Trạm Tay Gắp", "移载": "Chuyển Tải", "翻转": "Trạm Lật",
    "检测": "Kiểm Tra", "覆膜": "Dán Màng", "下料站": "Trạm Xuất Liệu",
    "料仓": "Silo (Kho)", "机械手": "Robot", "DataUpdata": "Cập Nhật Dữ Liệu",
    "初始化": "khởi tạo", "生产流程开始": "Bắt đầu quy trình sản xuất", "生产流程结束": "Kết thúc quy trình sản xuất",
    "流程开始": "bắt đầu quy trình",
    
    "运动到": "di chuyển tới", "移动到": "di chuyển tới", "到": "tới", "回": "về",
    "等待": "Chờ", "完成": "hoàn tất", "换料": "đổi vật liệu", "上料": "cấp liệu",
    "下料": "xuất liệu", "取料": "lấy liệu", "放料": "thả liệu", "拍照": "chụp ảnh",
    "破真空": "tắt hút chân không", "吸真空": "bật hút chân không", "真空": "chân không",
    "检查": "kiểm tra", "通知": "báo", "打开": "mở", "关闭": "đóng", "报错": "báo lỗi",
    "有料": "có vật liệu", "准备": "chuẩn bị", "下降": "hạ xuống", "上升": "nâng lên",
    
    "轴": "Trục ", "气缸": "Xi-lanh ", "升降": "Nâng hạ ", "吸嘴": "Vòi hút ",
    "钢环": "Vòng thép", "料盘": "Khay liệu", "顶针": "Kim đẩy", "流线": "Băng tải",
    "物料": "Vật liệu",
    
    "待机位": "Vị trí chờ", "上料位": "Vị trí cấp liệu", "取料位": "Vị trí lấy liệu",
    "放料位": "Vị trí thả liệu", "拍照位": "Vị trí chụp ảnh", "换料位": "Vị trí đổi liệu",
    "毫米": "mm", "不在": "không ở", "是否": "có/không", "超限": "quá giới hạn", "清料": "dọn liệu"
}

def translate_zh_vi(text):
    if not text: return text
    clean_text = text.replace('RegBit.', '').replace('DI.', '').replace('DO.', '').replace('AXIS.', '')
    
    res = clean_text
    sorted_keys = sorted(word_dict.keys(), key=len, reverse=True)
    for zh in sorted_keys:
        res = res.replace(zh, word_dict[zh] + " ")
        
    res = " ".join(res.split())
    if res == clean_text:
        return clean_text
    return res.strip()

def extract_method(content, method_name):
    pattern = r'protected\s+override\s+string\s+' + method_name + r'\s*\(\s*string\s+strStep\s*\)\s*\{'
    match = re.search(pattern, content)
    if not match: return None
    start_idx = match.end()
    open_brackets = 1
    end_idx = start_idx
    for i, char in enumerate(content[start_idx:], start_idx):
        if char == '{': open_brackets += 1
        elif char == '}':
            open_brackets -= 1
            if open_brackets == 0:
                end_idx = i
                break
    return content[start_idx:end_idx]

def extract_cases(method_body):
    cases = []
    case_pattern = re.compile(r'case\s+"([^"]+)":')
    matches = list(case_pattern.finditer(method_body))
    for i in range(len(matches)):
        case_name = matches[i].group(1)
        start_idx = matches[i].end()
        if i < len(matches) - 1:
            end_idx = matches[i+1].start()
        else:
            default_match = re.search(r'\bdefault\s*:', method_body[start_idx:])
            end_idx = start_idx + default_match.start() if default_match else len(method_body)
        cases.append({"step": case_name, "code": method_body[start_idx:end_idx].strip()})
    return cases

def parse_semantic_actions(code):
    actions = []
    
    for m in re.finditer(r'WaitRegBit\s*\(\s*RegBit\.([^,]+),\s*([^,\)]+)', code):
        trans = translate_zh_vi(m.group(1))
        actions.append(f"⏳ ĐỢI TÍN HIỆU CỜ: Chờ [{trans}] == {m.group(2).strip()}")
        
    for m in re.finditer(r'WaitDI\s*\(\s*DI\.([^,]+),\s*([^,\)]+)', code):
        trans = translate_zh_vi(m.group(1))
        actions.append(f"⏳ ĐỢI PHẦN CỨNG (DI): Chờ [{trans}] == {m.group(2).strip()}")
        
    for m in re.finditer(r'DO\.([^.]+)\.Set\s*\(\s*([^)]+)\)', code):
        trans = translate_zh_vi(m.group(1))
        val = m.group(2).strip()
        actions.append(f"🔌 ĐIỀU KHIỂN IO (DO): Bật/Tắt [{trans}] = {val}")

    for m in re.finditer(r'AxisSingleTo\s*\([^,]+,\s*AXIS\.([^,]+)', code):
        trans = translate_zh_vi(m.group(1))
        actions.append(f"🚀 DI CHUYỂN TRỤC: Yêu cầu [{trans}] chạy.")
        
    for m in re.finditer(r'WaitAxis\s*\(\s*AXIS\.([^,\)]+)', code):
        trans = translate_zh_vi(m.group(1))
        actions.append(f"⏳ ĐỢI TRỤC: Chờ [{trans}] tới nơi.")
        
    for m in re.finditer(r'RegBit\.([^.]+)\.Set\s*\(\s*([^)]+)\)', code):
        trans = translate_zh_vi(m.group(1))
        val = m.group(2).strip()
        actions.append(f"🚩 XUẤT TÍN HIỆU CỜ: Báo cáo [{trans}] = {val}")
        
    for m in re.finditer(r'return\s+"([^"]+)"', code):
        step_vi = translate_zh_vi(m.group(1))
        actions.append(f"⏭ BƯỚC TIẾP THEO LÀ: [{step_vi}]")

    return actions

doc = Document()
doc.add_heading('Văn Bản Lưu Trình (Tiếng Việt Toàn Diện)', 0)

for filename in os.listdir(station_dir):
    if not filename.endswith(".cs"): continue
    filepath = os.path.join(station_dir, filename)
    with open(filepath, 'r', encoding='utf-8-sig') as f:
        try: content = f.read()
        except UnicodeDecodeError:
            with open(filepath, 'r', encoding='gbk') as f2: content = f2.read()
            
    produce_body = extract_method(content, "ProduceRun")
    if not produce_body: continue
    
    st_name = translate_zh_vi(filename.replace(".cs", ""))
    doc.add_heading(f'🏭 TRẠM: {st_name}', level=1)
    
    cases = extract_cases(produce_body)
    for c in cases:
        step_vi = translate_zh_vi(c['step'])
        
        p = doc.add_paragraph()
        p.add_run(f"■ BƯỚC: [{step_vi}]").bold = True
        
        actions = parse_semantic_actions(c['code'])
        if actions:
            for act in actions:
                p_act = doc.add_paragraph(f"    {act}")
                if "⏳" in act: p_act.runs[0].font.color.rgb = RGBColor(200, 100, 0)
                elif "🚀" in act or "🔌" in act: p_act.runs[0].font.color.rgb = RGBColor(0, 100, 200)
                elif "🚩" in act: p_act.runs[0].font.color.rgb = RGBColor(200, 0, 200)
                elif "⏭" in act: p_act.runs[0].font.color.rgb = RGBColor(0, 150, 0)
        else:
            doc.add_paragraph("    (Chỉ xử lý tính toán nội bộ / Không có thao tác phần cứng/cờ báo hiệu)")

doc.save(output_file)
print(f"Successfully created {output_file}")
