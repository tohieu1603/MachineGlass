import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

station_dir = r"h:\水晶光电AOI程序1002\Autoframe\Station"
output_file = r"h:\水晶光电AOI程序1002\Station_Workflows_Detailed.docx"

def extract_method(content, method_name):
    pattern = r'protected\s+override\s+string\s+' + method_name + r'\s*\(\s*string\s+strStep\s*\)\s*\{'
    match = re.search(pattern, content)
    if not match:
        return None
    
    start_idx = match.end()
    open_brackets = 1
    end_idx = start_idx
    
    for i, char in enumerate(content[start_idx:], start_idx):
        if char == '{':
            open_brackets += 1
        elif char == '}':
            open_brackets -= 1
            if open_brackets == 0:
                end_idx = i
                break
                
    return content[start_idx:end_idx]

def extract_cases(method_body):
    cases = []
    # Match `case "StepName":`
    case_pattern = re.compile(r'case\s+"([^"]+)":')
    matches = list(case_pattern.finditer(method_body))
    
    for i in range(len(matches)):
        case_name = matches[i].group(1)
        start_idx = matches[i].end()
        
        # Determine end of this case block (either start of next case, or 'default:', or end of switch)
        if i < len(matches) - 1:
            end_idx = matches[i+1].start()
        else:
            default_match = re.search(r'\bdefault\s*:', method_body[start_idx:])
            if default_match:
                end_idx = start_idx + default_match.start()
            else:
                end_idx = len(method_body)
                
        case_body = method_body[start_idx:end_idx].strip()
        
        # Find all possible return paths
        return_matches = re.findall(r'return\s+"([^"]+)"', case_body)
        returns = list(set(return_matches)) # unique returns
        
        # Find comments for description
        comments = re.findall(r'//(.*)', case_body)
        description = "\n".join([c.strip() for c in comments if c.strip()])
        
        cases.append({
            "step": case_name,
            "returns": returns,
            "description": description,
            "code": case_body
        })
    return cases

doc = Document()
doc.add_heading('Lưu Trình Các Trạm - Chi Tiết (Detailed Workflows)', 0)

for filename in os.listdir(station_dir):
    if filename.endswith(".cs"):
        filepath = os.path.join(station_dir, filename)
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            try:
                content = f.read()
            except UnicodeDecodeError:
                with open(filepath, 'r', encoding='gbk') as f2:
                    content = f2.read()
        
        init_body = extract_method(content, "InitRun")
        produce_body = extract_method(content, "ProduceRun")
        
        if init_body or produce_body:
            doc.add_heading(f'Trạm: {filename.replace(".cs", "")}', level=1)
            
            for flow_name, body in [("1. Lưu trình Khởi tạo (InitRun)", init_body), ("2. Lưu trình Sản xuất (ProduceRun)", produce_body)]:
                if body:
                    doc.add_heading(flow_name, level=2)
                    cases = extract_cases(body)
                    if not cases:
                        doc.add_paragraph("Không có các bước cụ thể.")
                        continue
                    
                    for c in cases:
                        p = doc.add_paragraph()
                        # Tên bước
                        p.add_run(f"■ Bước: ").bold = True
                        p.add_run(f"[{c['step']}]\n").font.color.rgb = RGBColor(0, 0, 255)
                        
                        # Hướng rẽ
                        p.add_run(f"  ➔ Sẽ rẽ nhánh tới:\n").bold = True
                        if c['returns']:
                            for ret in c['returns']:
                                p.add_run(f"      - [{ret}]\n").font.color.rgb = RGBColor(0, 128, 0)
                        else:
                            p.add_run(f"      - (Có điều kiện ngầm / Kết thúc / Trả về null)\n").font.color.rgb = RGBColor(255, 0, 0)
                        
                        # Ghi chú
                        if c['description']:
                            p.add_run(f"  📝 Ghi chú trong code:\n").bold = True
                            p.add_run(f"      {c['description']}\n").font.color.rgb = RGBColor(128, 128, 128)
                        
                        # Code
                        doc.add_paragraph(f"  {'-'*40}\n  Mã nguồn gốc của bước này:")
                        code_para = doc.add_paragraph(c['code'].strip())
                        # Giảm size font code cho gọn
                        for run in code_para.runs:
                            run.font.name = 'Courier New'
                            run.font.size = Pt(9)
                        
                        doc.add_paragraph(f"  {'-'*40}\n")
        else:
            print(f"Skipping {filename}: No flows found.")

doc.save(output_file)
print(f"Successfully created {output_file}")
