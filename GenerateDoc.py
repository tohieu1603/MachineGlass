import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

station_dir = r"h:\水晶光电AOI程序1002\Autoframe\Station"
output_file = r"h:\水晶光电AOI程序1002\Station_Workflows.docx"

def extract_producerun(content):
    # Find the start of ProduceRun method
    match = re.search(r'protected\s+override\s+string\s+ProduceRun\s*\(\s*string\s+strStep\s*\)\s*\{', content)
    if not match:
        return None
    
    start_idx = match.end()
    open_brackets = 1
    end_idx = start_idx
    
    for i, char in enumerate(content[start_idx:], start_idx):
        if char == '{':
            open_brackets += 1
        elif char == '}':
            open_brackets -= 1
            if open_brackets == 0:
                end_idx = i
                break
                
    return content[start_idx:end_idx]

def extract_cases(produce_run_body):
    cases = []
    # Split by `case "StepName":`
    case_blocks = re.split(r'\bcase\s+"([^"]+)":', produce_run_body)
    
    if len(case_blocks) > 1:
        for i in range(1, len(case_blocks), 2):
            case_name = case_blocks[i]
            case_body = case_blocks[i+1]
            
            # extract the return string if any (first return statement)
            return_match = re.search(r'return\s+"([^"]+)"', case_body)
            next_step = return_match.group(1) if return_match else "Có điều kiện (Conditional) / Null"
            
            # clean up the body code snippet
            code_snippet = case_body.split('default:')[0]  # Avoid bleeding into default
            
            cases.append({
                "step": case_name,
                "next": next_step,
                "code": "case \"" + case_name + "\":\n" + code_snippet.strip()
            })
    return cases

doc = Document()
doc.add_heading('Lưu Trình Các Trạm (ProduceRun Workflow)', 0)

for filename in os.listdir(station_dir):
    if filename.endswith(".cs"):
        filepath = os.path.join(station_dir, filename)
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            try:
                content = f.read()
            except UnicodeDecodeError:
                with open(filepath, 'r', encoding='gbk') as f2:
                    content = f2.read()
        
        produce_run_body = extract_producerun(content)
        if produce_run_body:
            doc.add_heading(f'Trạm: {filename.replace(".cs", "")}', level=1)
            
            cases = extract_cases(produce_run_body)
            if not cases:
                doc.add_paragraph("Không tìm thấy các bước (cases) rõ ràng trong ProduceRun.")
                continue
                
            for c in cases:
                p = doc.add_paragraph()
                p.add_run(f"Bước: ").bold = True
                p.add_run(f"[{c['step']}]").font.color.rgb = RGBColor(0, 0, 255)
                
                p.add_run(f"  ➔ Tiếp theo: ").bold = True
                p.add_run(f"[{c['next']}]").font.color.rgb = RGBColor(0, 128, 0)
                
                # Add code block
                # Only keep the relevant lines for the snippet
                lines = c['code'].split('\n')
                clean_lines = []
                for line in lines:
                    if "case " in line and line != lines[0]:
                        break # Stop if we accidentally captured the next case string via regex failure
                    clean_lines.append(line)
                    if "return " in line:
                        break # Stop after the first return
                
                code_para = doc.add_paragraph('\n'.join(clean_lines))
        else:
            print(f"Skipping {filename}: No ProduceRun found.")

doc.save(output_file)
print(f"Successfully created {output_file}")
